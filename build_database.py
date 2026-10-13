import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

# Load verified student data
with open('all_students.json', encoding='utf-8') as f:
    pages = json.load(f)

wb = openpyxl.Workbook()

# Setup sheets
ws_search = wb.active
ws_search.title = "🔍 ÖĞRENCİ SORGULAMA"
ws_custom = wb.create_sheet(title="📝 ÖZEL LİSTE OLUŞTUR")
ws_classlist = wb.create_sheet(title="📄 SINIF LİSTESİ AL")
ws_db = wb.create_sheet(title="📋 ÖĞRENCİ VERİTABANI")
ws_classes = wb.create_sheet(title="📊 SINIFLAR VE İSTATİSTİK")

# Colors
DARK_NAVY = "1B365D"
PRIMARY_BLUE = "205493"
ACCENT_BLUE = "0078D4"
LIGHT_BG = "F8FAFC"
INPUT_BG = "FEF9C3"
SUCCESS_BG = "DEF7EC"
SUCCESS_TEXT = "03543F"
LABEL_BG = "E2E8F0"
WHITE = "FFFFFF"
BORDER_COLOR = "CBD5E1"
MUTED_TEXT = "475569"

font_main_title = Font(name="Segoe UI", size=14, bold=True, color=WHITE)
font_sub_title = Font(name="Segoe UI", size=10, bold=False, color="E2E8F0")
font_sec_header = Font(name="Segoe UI", size=11, bold=True, color=WHITE)
font_label = Font(name="Segoe UI", size=10, bold=True, color="334155")
font_value = Font(name="Segoe UI", size=11, bold=True, color="0F172A")
font_input = Font(name="Segoe UI", size=11.5, bold=True, color="1E3A8A")
font_tbl_hdr = Font(name="Segoe UI", size=10, bold=True, color=WHITE)
font_tbl_row = Font(name="Segoe UI", size=9.5, bold=False, color="1E293B")
font_tbl_bold = Font(name="Segoe UI", size=9.5, bold=True, color="1E293B")
font_hint = Font(name="Segoe UI", size=9, italic=True, color="64748B")

fill_banner = PatternFill(start_color=DARK_NAVY, end_color=DARK_NAVY, fill_type="solid")
fill_hdr_sec = PatternFill(start_color=PRIMARY_BLUE, end_color=PRIMARY_BLUE, fill_type="solid")
fill_accent = PatternFill(start_color=ACCENT_BLUE, end_color=ACCENT_BLUE, fill_type="solid")
fill_input = PatternFill(start_color=INPUT_BG, end_color=INPUT_BG, fill_type="solid")
fill_label = PatternFill(start_color=LABEL_BG, end_color=LABEL_BG, fill_type="solid")
fill_val = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")
fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
fill_success = PatternFill(start_color=SUCCESS_BG, end_color=SUCCESS_BG, fill_type="solid")
fill_selected = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

thin_side = Side(border_style="thin", color=BORDER_COLOR)
thick_side = Side(border_style="medium", color=PRIMARY_BLUE)
border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
border_input_box = Border(left=thick_side, right=thick_side, top=thick_side, bottom=thick_side)

# ====================================================
# 1. DATABASE SHEET ("📋 ÖĞRENCİ VERİTABANI")
# ====================================================
headers_db = [
    "Sıra No", "Öğrenci No", "Adı", "Soyadı", "Adı Soyadı",
    "Sınıf Seviyesi", "Şube", "Sınıfı / Şubesi", "Cinsiyeti",
    "Sınıf Rehber Öğretmeni", "Müdür Yardımcısı", "Sınıftaki Görevi",
    "Arama Etiketi", "Sınıf Filtre Sırası"
]

ws_db.row_dimensions[1].height = 26
for col_idx, h in enumerate(headers_db, 1):
    c = ws_db.cell(row=1, column=col_idx, value=h)
    c.font = font_tbl_hdr
    c.fill = fill_hdr_sec
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border_cell

db_row = 2
global_seq = 1

for p in pages:
    c_level = p["class_level"]
    c_branch = p["branch"]
    c_full = p["full_class"]
    c_teacher = p["teacher"]
    c_asst = p["assistant_principal"]
    pres = p["president"]
    v_pres = p["vice_president"]

    for s in p["students"]:
        s_no = s[1]
        first_name = s[2].strip()
        last_name = s[3].strip()
        full_name = f"{first_name} {last_name}"
        gender = s[4].strip()

        role = "Öğrenci"
        if pres and full_name == pres:
            role = "Sınıf Başkanı ⭐"
        elif v_pres and full_name == v_pres:
            role = "Sınıf Başkan Yrd. 🌟"

        search_label = f"{s_no} - {full_name} ({c_full})"
        class_helper_formula = f'=IF(H{db_row}=\'📄 SINIF LİSTESİ AL\'!$D$3, COUNTIF(\'📋 ÖĞRENCİ VERİTABANI\'!$H$2:H{db_row}, \'📄 SINIF LİSTESİ AL\'!$D$3), "")'

        row_vals = [
            global_seq, s_no, first_name, last_name, full_name,
            c_level, c_branch, c_full, gender,
            c_teacher if c_teacher else "-",
            c_asst if c_asst else "BARIŞ BAYRAK",
            role, search_label, class_helper_formula
        ]

        ws_db.row_dimensions[db_row].height = 19
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws_db.cell(row=db_row, column=col_idx, value=val)
            cell.font = font_tbl_row
            cell.border = border_cell
            if db_row % 2 == 1:
                cell.fill = fill_zebra
            
            if col_idx in [1, 2, 6, 7, 9, 14]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 12:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if "Başkan" in str(val):
                    cell.font = font_tbl_bold
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        global_seq += 1
        db_row += 1

last_db_row = db_row - 1

for col in ws_db.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_db.column_dimensions[col_letter].width = max(max_len + 3, 11)

ws_db.column_dimensions['N'].width = 16
ws_db.auto_filter.ref = f"A1:M{last_db_row}"
ws_db.freeze_panes = "A2"
ws_db.views.sheetView[0].showGridLines = True

# ====================================================
# 2. SEARCH SHEET ("🔍 ÖĞRENCİ SORGULAMA")
# ====================================================
ws_search.views.sheetView[0].showGridLines = True
ws_search.column_dimensions['A'].width = 3
ws_search.column_dimensions['B'].width = 24
ws_search.column_dimensions['C'].width = 30
ws_search.column_dimensions['D'].width = 24
ws_search.column_dimensions['E'].width = 34
ws_search.column_dimensions['F'].width = 16
ws_search.column_dimensions['G'].width = 16
ws_search.column_dimensions['H'].width = 4

ws_search.row_dimensions[2].height = 30
ws_search.merge_cells("B2:G2")
b1 = ws_search["B2"]
b1.value = "🏫 T.C. İZMİR VALİLİĞİ / TORBALI İLÇE MİLLÎ EĞİTİM MÜDÜRLÜĞÜ"
b1.font = font_main_title
b1.fill = fill_banner
b1.alignment = Alignment(horizontal="center", vertical="center")

ws_search.row_dimensions[3].height = 22
ws_search.merge_cells("B3:G3")
b2 = ws_search["B3"]
b2.value = "ŞEHİT KADİR ALTUNTAŞ CUMHURİYET ORTAOKULU - ÖĞRENCİ BİLGİ SORGULAMA PANELİ"
b2.font = font_sub_title
b2.fill = fill_banner
b2.alignment = Alignment(horizontal="center", vertical="center")

for r in [2, 3]:
    for c in range(2, 8):
        ws_search.cell(row=r, column=c).fill = fill_banner

ws_search.row_dimensions[5].height = 24
ws_search.merge_cells("B5:G5")
sh = ws_search["B5"]
sh.value = "🔎 TEK ÖĞRENCİ ARAMA KUTUSU"
sh.font = font_sec_header
sh.fill = fill_hdr_sec
sh.alignment = Alignment(horizontal="left", vertical="center", indent=1)

ws_search.row_dimensions[6].height = 28
ws_search["B6"] = "Öğrenci No veya Ad Soyad:"
ws_search["B6"].font = font_label
ws_search["B6"].fill = fill_label
ws_search["B6"].alignment = Alignment(horizontal="right", vertical="center")
ws_search["B6"].border = border_cell

ws_search["C6"] = 5015
ws_search["C6"].font = font_input
ws_search["C6"].fill = fill_input
ws_search["C6"].alignment = Alignment(horizontal="center", vertical="center")
ws_search["C6"].border = border_input_box

ws_search.merge_cells("D6:G6")
ws_search["D6"] = "👉 Doğrudan numara (5015, 6011, 239, 10...) veya isim yazıp Enter'a basabilirsiniz."
ws_search["D6"].font = font_hint
ws_search["D6"].alignment = Alignment(horizontal="left", vertical="center")
for c in range(4, 8):
    ws_search.cell(row=6, column=c).border = border_cell

ws_search.row_dimensions[7].height = 28
ws_search["B7"] = "Listeden Öğrenci Seç:"
ws_search["B7"].font = font_label
ws_search["B7"].fill = fill_label
ws_search["B7"].alignment = Alignment(horizontal="right", vertical="center")
ws_search["B7"].border = border_cell

ws_search["C7"] = ""
ws_search["C7"].font = font_input
ws_search["C7"].fill = fill_input
ws_search["C7"].alignment = Alignment(horizontal="center", vertical="center")
ws_search["C7"].border = border_input_box

ws_search.merge_cells("D7:G7")
ws_search["D7"] = "👉 Veya tüm okul listesinden açılır menü okuyla öğrenci seçebilirsiniz."
ws_search["D7"].font = font_hint
ws_search["D7"].alignment = Alignment(horizontal="left", vertical="center")
for c in range(4, 8):
    ws_search.cell(row=7, column=c).border = border_cell

dv = DataValidation(type="list", formula1=f"'📋 ÖĞRENCİ VERİTABANI'!$E$2:$E${last_db_row}", allow_blank=True)
dv.error = 'Listeden geçerli bir öğrenci seçiniz veya kutuyu boş bırakınız.'
dv.errorTitle = 'Geçersiz Seçim'
dv.prompt = 'Öğrenci adı seçmek için oka tıklayınız'
dv.promptTitle = 'Öğrenci Listesi'
ws_search.add_data_validation(dv)
dv.add("C7")

match_formula = (
    '=IF(AND(ISBLANK(C6), ISBLANK(C7)), "", '
    'IF(NOT(ISBLANK(C6)), '
    'IFERROR(MATCH(IF(ISNUMBER(C6), C6, VALUE(C6)), \'📋 ÖĞRENCİ VERİTABANI\'!$B$2:$B$498, 0), '
    'IFERROR(MATCH(C6, \'📋 ÖĞRENCİ VERİTABANI\'!$E$2:$E$498, 0), '
    'IFERROR(MATCH("*"&C6&"*", \'📋 ÖĞRENCİ VERİTABANI\'!$E$2:$E$498, 0), 0))), '
    'IFERROR(MATCH(C7, \'📋 ÖĞRENCİ VERİTABANI\'!$E$2:$E$498, 0), '
    'IFERROR(MATCH("*"&C7&"*", \'📋 ÖĞRENCİ VERİTABANI\'!$E$2:$E$498, 0), 0))))'
)
ws_search["H6"] = match_formula
ws_search["H6"].font = Font(name="Segoe UI", size=7, color="CBD5E1")

ws_search.row_dimensions[8].height = 24
ws_search.merge_cells("B8:G8")
st = ws_search["B8"]
st.value = '=IF(H6="", "ℹ️ Lütfen yukarıdaki arama kutusuna numara/isim yazınız veya listeden bir öğrenci seçiniz.", IF(H6=0, "⚠️ Aradığınız kriterde öğrenci bulunamadı. Lütfen numarayı veya adı kontrol ediniz.", "✅ Öğrenci kaydı başarıyla bulundu ve bilgileri aşağıda listelendi."))'
st.font = Font(name="Segoe UI", size=9.5, bold=True, color=SUCCESS_TEXT)
st.fill = fill_success
st.alignment = Alignment(horizontal="center", vertical="center")
for c in range(2, 8):
    ws_search.cell(row=8, column=c).fill = fill_success
    ws_search.cell(row=8, column=c).border = border_cell

ws_search.row_dimensions[10].height = 26
ws_search.merge_cells("B10:G10")
ct = ws_search["B10"]
ct.value = "🎓 ÖĞRENCİ BİLGİ VE KİMLİK KARTI"
ct.font = font_sec_header
ct.fill = fill_banner
ct.alignment = Alignment(horizontal="center", vertical="center")
for c in range(2, 8):
    ws_search.cell(row=10, column=c).fill = fill_banner

card_rows = [
    (11, "Okul Numarası:", "$B", "Sınıfı / Şubesi:", "$H"),
    (12, "Adı:", "$C", "Sınıf Seviyesi:", "$F"),
    (13, "Soyadı:", "$D", "Şubesi:", "$G"),
    (14, "Adı Soyadı:", "$E", "Cinsiyeti:", "$I"),
    (15, "Sınıf Rehber Öğretmeni:", "$J", "Sınıf Müdür Yrd.:", "$K"),
    (16, "Sınıftaki Görevi:", "$L", "Okul Kayıt Durumu:", "Aktif Öğrenci (2026-2027)")
]

for r_idx, lbl1, col1, lbl2, col2 in card_rows:
    ws_search.row_dimensions[r_idx].height = 25
    c_l1 = ws_search.cell(row=r_idx, column=2, value=lbl1)
    c_l1.font = font_label
    c_l1.fill = fill_label
    c_l1.alignment = Alignment(horizontal="right", vertical="center")
    c_l1.border = border_cell
    
    c_v1 = ws_search.cell(row=r_idx, column=3)
    c_v1.value = f'=IF(H6="", "-", IF(H6=0, "Bulunamadı", INDEX(\'📋 ÖĞRENCİ VERİTABANI\'!{col1}$2:{col1}$498, H6)))'
    c_v1.font = font_value
    c_v1.fill = fill_val
    c_v1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c_v1.border = border_cell
    
    c_l2 = ws_search.cell(row=r_idx, column=4, value=lbl2)
    c_l2.font = font_label
    c_l2.fill = fill_label
    c_l2.alignment = Alignment(horizontal="right", vertical="center")
    c_l2.border = border_cell
    
    ws_search.merge_cells(start_row=r_idx, start_column=5, end_row=r_idx, end_column=7)
    c_v2 = ws_search.cell(row=r_idx, column=5)
    if col2.startswith("$"):
        c_v2.value = f'=IF(H6="", "-", IF(H6=0, "Bulunamadı", INDEX(\'📋 ÖĞRENCİ VERİTABANI\'!{col2}$2:{col2}$498, H6)))'
    else:
        c_v2.value = f'=IF(H6="", "-", IF(H6=0, "Bulunamadı", "{col2}"))'
    c_v2.font = font_value
    c_v2.fill = fill_val
    c_v2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for c in range(5, 8):
        ws_search.cell(row=r_idx, column=c).border = border_cell

ws_search.row_dimensions[18].height = 24
ws_search.merge_cells("B18:G18")
cls_hdr = ws_search["B18"]
cls_hdr.value = "📊 BULUNAN ÖĞRENCİNİN SINIF MEVCUT BİLGİLERİ"
cls_hdr.font = font_sec_header
cls_hdr.fill = fill_hdr_sec
cls_hdr.alignment = Alignment(horizontal="center", vertical="center")
for c in range(2, 8):
    ws_search.cell(row=18, column=c).fill = fill_hdr_sec

ws_search.row_dimensions[19].height = 24
ws_search["B19"] = "Sınıf Toplam Mevcudu:"
ws_search["B19"].font = font_label
ws_search["B19"].fill = fill_label
ws_search["B19"].alignment = Alignment(horizontal="right", vertical="center")
ws_search["B19"].border = border_cell

ws_search["C19"] = '=IF(H6="", "-", IF(H6=0, "-", COUNTIF(\'📋 ÖĞRENCİ VERİTABANI\'!$H$2:$H$498, INDEX(\'📋 ÖĞRENCİ VERİTABANI\'!$H$2:$H$498, H6)) & " Öğrenci"))'
ws_search["C19"].font = font_value
ws_search["C19"].fill = fill_val
ws_search["C19"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws_search["C19"].border = border_cell

ws_search["D19"] = "Kız / Erkek Dağılımı:"
ws_search["D19"].font = font_label
ws_search["D19"].fill = fill_label
ws_search["D19"].alignment = Alignment(horizontal="right", vertical="center")
ws_search["D19"].border = border_cell

ws_search.merge_cells("E19:G19")
ws_search["E19"] = '=IF(H6="", "-", IF(H6=0, "-", COUNTIFS(\'📋 ÖĞRENCİ VERİTABANI\'!$H$2:$H$498, INDEX(\'📋 ÖĞRENCİ VERİTABANI\'!$H$2:$H$498, H6), \'📋 ÖĞRENCİ VERİTABANI\'!$I$2:$I$498, "Kız") & " Kız  |  " & COUNTIFS(\'📋 ÖĞRENCİ VERİTABANI\'!$H$2:$H$498, INDEX(\'📋 ÖĞRENCİ VERİTABANI\'!$H$2:$H$498, H6), \'📋 ÖĞRENCİ VERİTABANI\'!$I$2:$I$498, "Erkek") & " Erkek"))'
ws_search["E19"].font = font_value
ws_search["E19"].fill = fill_val
ws_search["E19"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
for c in range(5, 8):
    ws_search.cell(row=19, column=c).border = border_cell

# ====================================================
# 3. DUAL-INPUT CUSTOM LIST ("📝 ÖZEL LİSTE OLUŞTUR")
# ====================================================
ws_custom.views.sheetView[0].showGridLines = True

ws_custom.column_dimensions['A'].width = 7   # S.No
ws_custom.column_dimensions['B'].width = 25  # GİRİŞ ALANI (No veya İsim)
ws_custom.column_dimensions['C'].width = 14  # Öğrenci No (Otomatik)
ws_custom.column_dimensions['D'].width = 18  # Adı
ws_custom.column_dimensions['E'].width = 18  # Soyadı
ws_custom.column_dimensions['F'].width = 24  # Adı Soyadı
ws_custom.column_dimensions['G'].width = 22  # Sınıfı / Şubesi
ws_custom.column_dimensions['H'].width = 12  # Cinsiyeti
ws_custom.column_dimensions['I'].width = 24  # Sınıf Rehber Öğretmeni
ws_custom.column_dimensions['J'].width = 18  # Sınıftaki Görevi

ws_custom.row_dimensions[1].height = 28
ws_custom.merge_cells("A1:J1")
c_b1 = ws_custom["A1"]
c_b1.value = "🎯 ÖZEL ÖĞRENCİ LİSTESİ OLUŞTURUCU (NUMARA VEYA İSİM GİRİŞLİ)"
c_b1.font = font_main_title
c_b1.fill = fill_banner
c_b1.alignment = Alignment(horizontal="center", vertical="center")
for c in range(1, 11):
    ws_custom.cell(row=1, column=c).fill = fill_banner

ws_custom.row_dimensions[2].height = 26
ws_custom.merge_cells("A2:E2")
c_info = ws_custom["A2"]
c_info.value = "💡 B sütununa ister OKUL NUMARASI, ister İSİM/SOYİSİM yazınız ya da yapıştırınız. Tüm bilgiler sağa otomatik dolar."
c_info.font = Font(name="Segoe UI", size=9.5, italic=False, bold=True, color=DARK_NAVY)
c_info.fill = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
c_info.alignment = Alignment(horizontal="left", vertical="center", indent=1)
for c in range(1, 6):
    ws_custom.cell(row=2, column=c).fill = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
    ws_custom.cell(row=2, column=c).border = border_cell

ws_custom["F2"] = "Toplam Öğrenci:"
ws_custom["F2"].font = font_label
ws_custom["F2"].fill = fill_label
ws_custom["F2"].alignment = Alignment(horizontal="right", vertical="center")
ws_custom["F2"].border = border_cell

ws_custom["G2"] = '=COUNT(C4:C300) & " Öğrenci"'
ws_custom["G2"].font = font_value
ws_custom["G2"].fill = fill_val
ws_custom["G2"].alignment = Alignment(horizontal="center", vertical="center")
ws_custom["G2"].border = border_cell

ws_custom["H2"] = "Kız / Erkek:"
ws_custom["H2"].font = font_label
ws_custom["H2"].fill = fill_label
ws_custom["H2"].alignment = Alignment(horizontal="right", vertical="center")
ws_custom["H2"].border = border_cell

ws_custom.merge_cells("I2:J2")
ws_custom["I2"] = '=COUNTIF(H4:H300, "Kız") & " Kız  |  " & COUNTIF(H4:H300, "Erkek") & " Erkek"'
ws_custom["I2"].font = font_value
ws_custom["I2"].fill = fill_val
ws_custom["I2"].alignment = Alignment(horizontal="center", vertical="center")
for c in [9, 10]:
    ws_custom.cell(row=2, column=c).border = border_cell

custom_headers = [
    ("S.No", fill_hdr_sec),
    ("Giriş (No veya İsim) ✍️", PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")),
    ("Öğrenci No", fill_hdr_sec),
    ("Adı", fill_hdr_sec),
    ("Soyadı", fill_hdr_sec),
    ("Adı Soyadı", fill_hdr_sec),
    ("Sınıfı / Şubesi", fill_hdr_sec),
    ("Cinsiyeti", fill_hdr_sec),
    ("Sınıf Rehber Öğretmeni", fill_hdr_sec),
    ("Sınıftaki Görevi", fill_hdr_sec)
]

ws_custom.row_dimensions[3].height = 26
for col_idx, (h_text, h_fill) in enumerate(custom_headers, 1):
    c = ws_custom.cell(row=3, column=col_idx, value=h_text)
    c.font = font_tbl_hdr
    c.fill = h_fill
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = border_cell

for r in range(4, 301):
    ws_custom.row_dimensions[r].height = 20
    
    cA = ws_custom.cell(row=r, column=1)
    cA.value = f'=IF(C{r}="","",ROW()-3)'
    cA.font = font_tbl_bold
    cA.alignment = Alignment(horizontal="center", vertical="center")
    cA.border = border_cell
    
    cB = ws_custom.cell(row=r, column=2)
    cB.font = font_input
    cB.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    cB.border = border_cell
    cB.fill = fill_input
    
    cC = ws_custom.cell(row=r, column=3)
    cC.value = (
        f'=IF(ISBLANK($B{r}), "", '
        f'IFERROR(INDEX(\'📋 ÖĞRENCİ VERİTABANI\'!$B$2:$B$498, '
        f'IFERROR(MATCH(IF(ISNUMBER($B{r}), $B{r}, VALUE($B{r})), \'📋 ÖĞRENCİ VERİTABANI\'!$B$2:$B$498, 0), '
        f'IFERROR(MATCH($B{r}, \'📋 ÖĞRENCİ VERİTABANI\'!$E$2:$E$498, 0), '
        f'MATCH("*"&$B{r}&"*", \'📋 ÖĞRENCİ VERİTABANI\'!$E$2:$E$498, 0)))), '
        f'"Bulunamadı"))'
    )
    cC.font = font_tbl_bold
    cC.alignment = Alignment(horizontal="center", vertical="center")
    cC.border = border_cell
    
    cD = ws_custom.cell(row=r, column=4)
    cD.value = f'=IF(OR($C{r}="", $C{r}="Bulunamadı"), "", INDEX(\'📋 ÖĞRENCİ VERİTABANI\'!$C$2:$C$498, MATCH($C{r}, \'📋 ÖĞRENCİ VERİTABANI\'!$B$2:$B$498, 0)))'
    cD.font = font_tbl_row
    cD.alignment = Alignment(horizontal="left", vertical="center")
    cD.border = border_cell
    
    cE = ws_custom.cell(row=r, column=5)
    cE.value = f'=IF(OR($C{r}="", $C{r}="Bulunamadı"), "", INDEX(\'📋 ÖĞRENCİ VERİTABANI\'!$D$2:$D$498, MATCH($C{r}, \'📋 ÖĞRENCİ VERİTABANI\'!$B$2:$B$498, 0)))'
    cE.font = font_tbl_row
    cE.alignment = Alignment(horizontal="left", vertical="center")
    cE.border = border_cell
    
    cF = ws_custom.cell(row=r, column=6)
    cF.value = f'=IF(OR($C{r}="", $C{r}="Bulunamadı"), "", INDEX(\'📋 ÖĞRENCİ VERİTABANI\'!$E$2:$E$498, MATCH($C{r}, \'📋 ÖĞRENCİ VERİTABANI\'!$B$2:$B$498, 0)))'
    cF.font = font_tbl_bold
    cF.alignment = Alignment(horizontal="left", vertical="center")
    cF.border = border_cell
    
    cG = ws_custom.cell(row=r, column=7)
    cG.value = f'=IF(OR($C{r}="", $C{r}="Bulunamadı"), "", INDEX(\'📋 ÖĞRENCİ VERİTABANI\'!$H$2:$H$498, MATCH($C{r}, \'📋 ÖĞRENCİ VERİTABANI\'!$B$2:$B$498, 0)))'
    cG.font = font_tbl_row
    cG.alignment = Alignment(horizontal="left", vertical="center")
    cG.border = border_cell
    
    cH = ws_custom.cell(row=r, column=8)
    cH.value = f'=IF(OR($C{r}="", $C{r}="Bulunamadı"), "", INDEX(\'📋 ÖĞRENCİ VERİTABANI\'!$I$2:$I$498, MATCH($C{r}, \'📋 ÖĞRENCİ VERİTABANI\'!$B$2:$B$498, 0)))'
    cH.font = font_tbl_row
    cH.alignment = Alignment(horizontal="center", vertical="center")
    cH.border = border_cell
    
    cI = ws_custom.cell(row=r, column=9)
    cI.value = f'=IF(OR($C{r}="", $C{r}="Bulunamadı"), "", INDEX(\'📋 ÖĞRENCİ VERİTABANI\'!$J$2:$J$498, MATCH($C{r}, \'📋 ÖĞRENCİ VERİTABANI\'!$B$2:$B$498, 0)))'
    cI.font = font_tbl_row
    cI.alignment = Alignment(horizontal="left", vertical="center")
    cI.border = border_cell
    
    cJ = ws_custom.cell(row=r, column=10)
    cJ.value = f'=IF(OR($C{r}="", $C{r}="Bulunamadı"), "", INDEX(\'📋 ÖĞRENCİ VERİTABANI\'!$L$2:$L$498, MATCH($C{r}, \'📋 ÖĞRENCİ VERİTABANI\'!$B$2:$B$498, 0)))'
    cJ.font = font_tbl_row
    cJ.alignment = Alignment(horizontal="center", vertical="center")
    cJ.border = border_cell
    
    if r % 2 == 1:
        for c_idx in [1, 3, 4, 5, 6, 7, 8, 9, 10]:
            ws_custom.cell(row=r, column=c_idx).fill = fill_zebra

sample_entries = [
    5015,
    "AHMET TAHA ÖZ",
    239,
    "KUBİLAY ATEŞ",
    "SÜMEYYE İKBAL"
]
for idx, s_val in enumerate(sample_entries, 4):
    ws_custom.cell(row=idx, column=2, value=s_val)

ws_custom.freeze_panes = "A4"

# ====================================================
# 4. CLASS ROSTER SHEET ("📄 SINIF LİSTESİ AL")
# ====================================================
ws_classlist.views.sheetView[0].showGridLines = True

ws_classlist.column_dimensions['A'].width = 7   # S.No
ws_classlist.column_dimensions['B'].width = 14  # Öğrenci No
ws_classlist.column_dimensions['C'].width = 22  # Adı
ws_classlist.column_dimensions['D'].width = 20  # Soyadı
ws_classlist.column_dimensions['E'].width = 12  # Cinsiyeti
ws_classlist.column_dimensions['F'].width = 22  # Sınıftaki Görevi
ws_classlist.column_dimensions['G'].width = 18  # Seçim / Tıkla Ekle sütunu ✍️

# Header Banner
ws_classlist.row_dimensions[1].height = 24
ws_classlist.merge_cells("A1:G1")
cl_b1 = ws_classlist["A1"]
cl_b1.value = "T.C. İZMİR VALİLİĞİ - TORBALI İLÇE MİLLÎ EĞİTİM MÜDÜRLÜĞÜ"
cl_b1.font = Font(name="Segoe UI", size=11, bold=True, color=DARK_NAVY)
cl_b1.alignment = Alignment(horizontal="center", vertical="center")

ws_classlist.row_dimensions[2].height = 22
ws_classlist.merge_cells("A2:G2")
cl_b2 = ws_classlist["A2"]
cl_b2.value = "Örnek Ortaokulu Müdürlüğü"
cl_b2.font = Font(name="Segoe UI", size=12, bold=True, color=PRIMARY_BLUE)
cl_b2.alignment = Alignment(horizontal="center", vertical="center")

# Class Selection Dropdown Area (Row 3)
ws_classlist.row_dimensions[3].height = 30
ws_classlist.merge_cells("A3:C3")
cl_sel_lbl = ws_classlist["A3"]
cl_sel_lbl.value = "📋 Listesini Almak İstediğiniz Sınıfı Seçiniz:"
cl_sel_lbl.font = Font(name="Segoe UI", size=11, bold=True, color="1E3A8A")
cl_sel_lbl.alignment = Alignment(horizontal="right", vertical="center")
cl_sel_lbl.fill = fill_label
for c in range(1, 4):
    ws_classlist.cell(row=3, column=c).fill = fill_label
    ws_classlist.cell(row=3, column=c).border = border_cell

ws_classlist.merge_cells("D3:F3")
cl_sel = ws_classlist["D3"]
cl_sel.value = "5. Sınıf / A Şubesi"
cl_sel.font = Font(name="Segoe UI", size=12, bold=True, color="92400E")
cl_sel.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
cl_sel.alignment = Alignment(horizontal="center", vertical="center")
for c in range(4, 7):
    ws_classlist.cell(row=3, column=c).fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    ws_classlist.cell(row=3, column=c).border = border_input_box

ws_classlist["G3"] = "💡 Çift Tıkla Seç"
ws_classlist["G3"].font = Font(name="Segoe UI", size=9, bold=True, color=PRIMARY_BLUE)
ws_classlist["G3"].fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
ws_classlist["G3"].alignment = Alignment(horizontal="center", vertical="center")
ws_classlist["G3"].border = border_cell

dv_cls = DataValidation(type="list", formula1=f"'📊 SINIFLAR VE İSTATİSTİK'!$B$5:$B$34", allow_blank=False)
dv_cls.error = 'Lütfen listeden geçerli bir sınıf seçiniz.'
dv_cls.errorTitle = 'Geçersiz Sınıf'
dv_cls.prompt = 'Sınıf seçmek için oka tıklayınız'
dv_cls.promptTitle = 'Sınıf Seçimi'
ws_classlist.add_data_validation(dv_cls)
dv_cls.add("D3")

# Title of the Active Class List (Row 4)
ws_classlist.row_dimensions[4].height = 26
ws_classlist.merge_cells("A4:G4")
cl_title = ws_classlist["A4"]
cl_title.value = '=D3 & " Sınıf Listesi"'
cl_title.font = Font(name="Segoe UI", size=12, bold=True, color=WHITE)
cl_title.fill = fill_banner
cl_title.alignment = Alignment(horizontal="center", vertical="center")
for c in range(1, 8):
    ws_classlist.cell(row=4, column=c).fill = fill_banner

# Metadata block Rows 5-6
ws_classlist.row_dimensions[5].height = 22
ws_classlist["A5"] = "Sınıf Öğretmeni:"
ws_classlist["A5"].font = font_label
ws_classlist["A5"].alignment = Alignment(horizontal="right", vertical="center")

ws_classlist.merge_cells("B5:C5")
ws_classlist["B5"] = '=IFERROR(VLOOKUP(D3, \'📊 SINIFLAR VE İSTATİSTİK\'!$B$5:$H$34, 2, FALSE), "-")'
ws_classlist["B5"].font = font_value
ws_classlist["B5"].alignment = Alignment(horizontal="left", vertical="center")

ws_classlist["D5"] = "Sınıf Başkanı:"
ws_classlist["D5"].font = font_label
ws_classlist["D5"].alignment = Alignment(horizontal="right", vertical="center")

ws_classlist.merge_cells("E5:F5")
ws_classlist["E5"] = '=IFERROR(VLOOKUP(D3, \'📊 SINIFLAR VE İSTATİSTİK\'!$B$5:$H$34, 3, FALSE), "-")'
ws_classlist["E5"].font = font_value
ws_classlist["E5"].alignment = Alignment(horizontal="left", vertical="center")

ws_classlist["G5"] = ""
ws_classlist["G5"].border = border_cell
ws_classlist["G5"].fill = fill_zebra

ws_classlist.row_dimensions[6].height = 22
ws_classlist["A6"] = "Müdür Yardımcısı:"
ws_classlist["A6"].font = font_label
ws_classlist["A6"].alignment = Alignment(horizontal="right", vertical="center")

ws_classlist.merge_cells("B6:C6")
ws_classlist["B6"] = "BARIŞ BAYRAK"
ws_classlist["B6"].font = font_value
ws_classlist["B6"].alignment = Alignment(horizontal="left", vertical="center")

ws_classlist["D6"] = "Sınıf Başkan Yrd.:"
ws_classlist["D6"].font = font_label
ws_classlist["D6"].alignment = Alignment(horizontal="right", vertical="center")

ws_classlist.merge_cells("E6:F6")
ws_classlist["E6"] = '=IFERROR(VLOOKUP(D3, \'📊 SINIFLAR VE İSTATİSTİK\'!$B$5:$H$34, 4, FALSE), "-")'
ws_classlist["E6"].font = font_value
ws_classlist["E6"].alignment = Alignment(horizontal="left", vertical="center")

ws_classlist["G6"] = ""
ws_classlist["G6"].border = border_cell
ws_classlist["G6"].fill = fill_zebra

for r in [5, 6]:
    for c in range(1, 8):
        ws_classlist.cell(row=r, column=c).border = border_cell
        ws_classlist.cell(row=r, column=c).fill = fill_zebra

# Table Headers (Row 7)
roster_headers = [
    ("S.No", fill_hdr_sec),
    ("Öğrenci No", fill_hdr_sec),
    ("Adı", fill_hdr_sec),
    ("Soyadı", fill_hdr_sec),
    ("Cinsiyeti", fill_hdr_sec),
    ("Sınıftaki Görevi", fill_hdr_sec),
    ("Seçim / Ekle ✍️", PatternFill(start_color="047857", end_color="047857", fill_type="solid"))
]
ws_classlist.row_dimensions[7].height = 24
for col_idx, (h_text, h_fill) in enumerate(roster_headers, 1):
    c = ws_classlist.cell(row=7, column=col_idx, value=h_text)
    c.font = font_tbl_hdr
    c.fill = h_fill
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = border_cell

# Dropdown Data Validation for Column G (so users can also click the dropdown to select!)
dv_sel = DataValidation(type="list", formula1='"[X] SECILDI"', allow_blank=True)
dv_sel.error = 'Lütfen [X] SECILDI seçiniz veya boş bırakınız.'
dv_sel.errorTitle = 'Geçersiz Seçim'
ws_classlist.add_data_validation(dv_sel)

# Roster Rows: 8 to 45
for r in range(8, 46):
    ws_classlist.row_dimensions[r].height = 20
    seq_idx = r - 7
    
    cA = ws_classlist.cell(row=r, column=1)
    cA.value = f'=IF(B{r}="","", {seq_idx})'
    cA.font = font_tbl_bold
    cA.alignment = Alignment(horizontal="center", vertical="center")
    cA.border = border_cell
    
    cB = ws_classlist.cell(row=r, column=2)
    cB.value = f'=IFERROR(INDEX(\'📋 ÖĞRENCİ VERİTABANI\'!$B$2:$B$498, MATCH({seq_idx}, \'📋 ÖĞRENCİ VERİTABANI\'!$N$2:$N$498, 0)), "")'
    cB.font = font_tbl_bold
    cB.alignment = Alignment(horizontal="center", vertical="center")
    cB.border = border_cell
    
    cC = ws_classlist.cell(row=r, column=3)
    cC.value = f'=IFERROR(INDEX(\'📋 ÖĞRENCİ VERİTABANI\'!$C$2:$C$498, MATCH({seq_idx}, \'📋 ÖĞRENCİ VERİTABANI\'!$N$2:$N$498, 0)), "")'
    cC.font = font_tbl_row
    cC.alignment = Alignment(horizontal="left", vertical="center")
    cC.border = border_cell
    
    cD = ws_classlist.cell(row=r, column=4)
    cD.value = f'=IFERROR(INDEX(\'📋 ÖĞRENCİ VERİTABANI\'!$D$2:$D$498, MATCH({seq_idx}, \'📋 ÖĞRENCİ VERİTABANI\'!$N$2:$N$498, 0)), "")'
    cD.font = font_tbl_row
    cD.alignment = Alignment(horizontal="left", vertical="center")
    cD.border = border_cell
    
    cE = ws_classlist.cell(row=r, column=5)
    cE.value = f'=IFERROR(INDEX(\'📋 ÖĞRENCİ VERİTABANI\'!$I$2:$I$498, MATCH({seq_idx}, \'📋 ÖĞRENCİ VERİTABANI\'!$N$2:$N$498, 0)), "")'
    cE.font = font_tbl_row
    cE.alignment = Alignment(horizontal="center", vertical="center")
    cE.border = border_cell
    
    cF = ws_classlist.cell(row=r, column=6)
    cF.value = f'=IFERROR(INDEX(\'📋 ÖĞRENCİ VERİTABANI\'!$L$2:$L$498, MATCH({seq_idx}, \'📋 ÖĞRENCİ VERİTABANI\'!$N$2:$N$498, 0)), "")'
    cF.font = font_tbl_row
    cF.alignment = Alignment(horizontal="center", vertical="center")
    cF.border = border_cell
    
    cG = ws_classlist.cell(row=r, column=7)
    cG.font = Font(name="Segoe UI", size=9.5, bold=True, color="047857")
    cG.alignment = Alignment(horizontal="center", vertical="center")
    cG.border = border_cell
    dv_sel.add(f"G{r}")
    
    if r % 2 == 1:
        for c_idx in range(1, 8):
            ws_classlist.cell(row=r, column=c_idx).fill = fill_zebra

# Add Native Excel Conditional Formatting:
# If Col G is not empty ($G8<>""), highlight row A8:G45 in light green!
rule_selected = FormulaRule(formula=['$G8<>""'], stopIfTrue=False, fill=fill_selected)
ws_classlist.conditional_formatting.add("A8:G45", rule_selected)

# Footer Summary (Row 46)
ws_classlist.row_dimensions[46].height = 24
ws_classlist.merge_cells("A46:B46")
ws_classlist["A46"] = '= "Kız Öğrenci: " & COUNTIF(E8:E45, "Kız")'
ws_classlist["A46"].font = font_tbl_bold
ws_classlist["A46"].alignment = Alignment(horizontal="center", vertical="center")
ws_classlist["A46"].fill = fill_label
ws_classlist["B46"].fill = fill_label
ws_classlist["A46"].border = border_cell
ws_classlist["B46"].border = border_cell

ws_classlist.merge_cells("C46:D46")
ws_classlist["C46"] = '= "Erkek Öğrenci: " & COUNTIF(E8:E45, "Erkek")'
ws_classlist["C46"].font = font_tbl_bold
ws_classlist["C46"].alignment = Alignment(horizontal="center", vertical="center")
ws_classlist["C46"].fill = fill_label
ws_classlist["D46"].fill = fill_label
ws_classlist["C46"].border = border_cell
ws_classlist["D46"].border = border_cell

ws_classlist.merge_cells("E46:G46")
ws_classlist["E46"] = '= "Toplam Mevcut: " & COUNT(B8:B45)'
ws_classlist["E46"].font = Font(name="Segoe UI", size=10, bold=True, color=DARK_NAVY)
ws_classlist["E46"].alignment = Alignment(horizontal="center", vertical="center")
ws_classlist["E46"].fill = fill_label
ws_classlist["F46"].fill = fill_label
ws_classlist["G46"].fill = fill_label
ws_classlist["E46"].border = border_cell
ws_classlist["F46"].border = border_cell
ws_classlist["G46"].border = border_cell

# Signature section
ws_classlist.row_dimensions[48].height = 20
ws_classlist["B48"] = "Sınıf Rehber Öğretmeni"
ws_classlist["B48"].font = Font(name="Segoe UI", size=10, bold=True, color=DARK_NAVY)
ws_classlist["B48"].alignment = Alignment(horizontal="center", vertical="center")

ws_classlist["E48"] = "Okul Müdürü"
ws_classlist["E48"].font = Font(name="Segoe UI", size=10, bold=True, color=DARK_NAVY)
ws_classlist["E48"].alignment = Alignment(horizontal="center", vertical="center")

ws_classlist.row_dimensions[49].height = 20
ws_classlist["B49"] = "=B5"
ws_classlist["B49"].font = font_tbl_row
ws_classlist["B49"].alignment = Alignment(horizontal="center", vertical="center")

ws_classlist["E49"] = "İmza / Mühür"
ws_classlist["E49"].font = font_hint
ws_classlist["E49"].alignment = Alignment(horizontal="center", vertical="center")

# ====================================================
# 5. CLASSES & STATS SHEET ("📊 SINIFLAR VE İSTATİSTİK")
# ====================================================
ws_classes.views.sheetView[0].showGridLines = True
ws_classes.column_dimensions['A'].width = 8
ws_classes.column_dimensions['B'].width = 28
ws_classes.column_dimensions['C'].width = 28
ws_classes.column_dimensions['D'].width = 22
ws_classes.column_dimensions['E'].width = 22
ws_classes.column_dimensions['F'].width = 14
ws_classes.column_dimensions['G'].width = 14
ws_classes.column_dimensions['H'].width = 16

ws_classes.row_dimensions[2].height = 28
ws_classes.merge_cells("A2:H2")
c_banner = ws_classes["A2"]
c_banner.value = "📊 ŞEHİT KADİR ALTUNTAŞ CUMHURİYET ORTAOKULU - ŞUBE BİLGİLERİ VE MEVCUTLARI"
c_banner.font = Font(name="Segoe UI", size=13, bold=True, color=WHITE)
c_banner.fill = fill_banner
c_banner.alignment = Alignment(horizontal="center", vertical="center")
for col_idx in range(1, 9):
    ws_classes.cell(row=2, column=col_idx).fill = fill_banner

headers_classes = [
    "S.No", "Sınıf ve Şube", "Sınıf Rehber Öğretmeni",
    "Sınıf Başkanı", "Sınıf Başkan Yrd.",
    "Kız Öğrenci", "Erkek Öğrenci", "Toplam Mevcut"
]

ws_classes.row_dimensions[4].height = 24
for col_idx, h in enumerate(headers_classes, 1):
    cell = ws_classes.cell(row=4, column=col_idx, value=h)
    cell.font = font_tbl_hdr
    cell.fill = fill_hdr_sec
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border_cell

cls_row = 5
for idx, p in enumerate(pages, 1):
    c_full = p["full_class"]
    c_teacher = p["teacher"] if p["teacher"] else "-"
    pres = p["president"] if p["president"] else "-"
    v_pres = p["vice_president"] if p["vice_president"] else "-"
    
    f_girl = f'=COUNTIFS(\'📋 ÖĞRENCİ VERİTABANI\'!$H$2:$H$498, "{c_full}", \'📋 ÖĞRENCİ VERİTABANI\'!$I$2:$I$498, "Kız")'
    f_boy = f'=COUNTIFS(\'📋 ÖĞRENCİ VERİTABANI\'!$H$2:$H$498, "{c_full}", \'📋 ÖĞRENCİ VERİTABANI\'!$I$2:$I$498, "Erkek")'
    f_total = f'=F{cls_row}+G{cls_row}'
    
    ws_classes.row_dimensions[cls_row].height = 19
    ws_classes.cell(row=cls_row, column=1, value=idx)
    ws_classes.cell(row=cls_row, column=2, value=c_full)
    ws_classes.cell(row=cls_row, column=3, value=c_teacher)
    ws_classes.cell(row=cls_row, column=4, value=pres)
    ws_classes.cell(row=cls_row, column=5, value=v_pres)
    ws_classes.cell(row=cls_row, column=6, value=f_girl)
    ws_classes.cell(row=cls_row, column=7, value=f_boy)
    ws_classes.cell(row=cls_row, column=8, value=f_total)
    
    for c in range(1, 9):
        cell = ws_classes.cell(row=cls_row, column=c)
        cell.font = font_tbl_row
        cell.border = border_cell
        if cls_row % 2 == 1:
            cell.fill = fill_zebra
        if c in [1, 6, 7, 8]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if c == 8:
                cell.font = font_tbl_bold
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
            
    cls_row += 1

ws_classes.row_dimensions[cls_row].height = 24
ws_classes.merge_cells(start_row=cls_row, start_column=1, end_row=cls_row, end_column=5)
t_cell = ws_classes.cell(row=cls_row, column=1, value="GENEL OKUL TOPLAMI")
t_cell.font = Font(name="Segoe UI", size=10.5, bold=True, color=WHITE)
t_cell.fill = fill_banner
t_cell.alignment = Alignment(horizontal="center", vertical="center")

for c in range(1, 6):
    ws_classes.cell(row=cls_row, column=c).fill = fill_banner
    ws_classes.cell(row=cls_row, column=c).border = border_cell

ws_classes.cell(row=cls_row, column=6, value=f"=SUM(F5:F{cls_row-1})")
ws_classes.cell(row=cls_row, column=7, value=f"=SUM(G5:G{cls_row-1})")
ws_classes.cell(row=cls_row, column=8, value=f"=SUM(H5:H{cls_row-1})")

for c in [6, 7, 8]:
    cell = ws_classes.cell(row=cls_row, column=c)
    cell.font = Font(name="Segoe UI", size=11, bold=True, color=WHITE)
    cell.fill = fill_banner
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border_cell

desktop_xlsx = r"C:\Users\SKACO-02\Desktop\Ogrenci_Veritabani.xlsx"
local_xlsx = "Ogrenci_Veritabani.xlsx"

wb.save(desktop_xlsx)
wb.save(local_xlsx)
print("Base XLSX with conditional formatting and dropdowns generated successfully.")
